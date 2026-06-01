import httpx
import csv
import io
import os
from pathlib import Path
from typing import List, Dict, Any, Optional
from openpyxl import load_workbook

SHEET_URLS = {
    2026: "1jZPBx9QYFYwsBWMF7_QtwXM2Rjt4oJyKJ6zSPLlebu8",
}

DATA_DIR = Path(__file__).parent.parent / "data"

EXCEL_FILES = {
    2025: DATA_DIR / "2025_Ekstraliga_Stats.xlsx",
}

HITTING_CORE_COLS = [
    "Name", "Team", "G", "PA", "AB", "R", "H", "1B", "2B", "3B", "HR",
    "RBI", "AVG", "TB", "BB", "Kc", "Ks", "SO", "HBP", "SB", "CS",
    "SB%", "OBP", "SLG", "OPS", "ISO", "BABIP", "OPS+", "JOPS", "pJOPS",
    "K%", "BB%", "BB%-K%", "RC", "RC/PA", "ROE", "FC", "GDP",
    "AB/RSP", "H/RSP", "BA/RSP", "SWINGS", "BIP", "CON%", "P/PA",
]

PITCHING_CORE_COLS = [
    "Name", "Team", "G", "W", "L", "SV", "BS", "GS", "IP", "BF",
    "R", "ER", "ERA7", "ERA9", "WHIP", "FIP", "K", "Kc", "Ks",
    "H", "BB", "IBB", "K/BB", "K/9", "BB/9", "H/9", "HB", "BK",
    "WP", "HR", "GO", "AO", "FPS%", "E+A%", "STR%", "CSW%", "WHF%",
    "BB%", "K%",
]

FIELDING_CORE_COLS = [
    "Name", "Team", "G", "ERR", "PO", "A", "SBA", "CS", "CS%",
    "DP", "TP", "PB", "PKF", "PK", "FP",
    "Start P", "Start C", "Start 1B", "Start 2B", "Start 3B",
    "Start SS", "Start LF", "Start CF", "Start RF", "Start OF", "Start DH",
]


class SheetsService:

    # ─────────────────────────────────────────────
    # EXCEL READER (2025)
    # ─────────────────────────────────────────────

    def _read_excel_sheet(self, year: int, sheet_name: str, core_cols: List[str]) -> List[Dict[str, Any]]:
        path = EXCEL_FILES.get(year)
        if not path or not path.exists():
            raise FileNotFoundError(f"Excel file not found: {path}")

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet_name]

        rows = ws.iter_rows(values_only=True)
        headers = [str(h).strip() if h else "" for h in next(rows)]

        result = []
        for row in rows:
            if not row[0]:
                continue
            row_dict = dict(zip(headers, row))
            name = str(row_dict.get("Name", "")).strip()
            if not name or name == "Name":
                continue

            # Clean up values
            cleaned = {}
            for col in core_cols:
                val = row_dict.get(col, None)
                if val in ["#DIV/0!", "#VALUE!", "#REF!", "#N/A", "#NAME?", None]:
                    cleaned[col] = None
                elif isinstance(val, float):
                    # Round to 3 decimal places for readability
                    cleaned[col] = round(val, 3)
                else:
                    cleaned[col] = val
            result.append(cleaned)

        wb.close()
        return result

    # ─────────────────────────────────────────────
    # GOOGLE SHEETS CSV READER (2026)
    # ─────────────────────────────────────────────

    def _csv_url(self, sheet_id: str, gid: str = "0") -> str:
        return f"https://docs.google.com/spreadsheets/d/{sheet_id}/gviz/tq?tqx=out:csv&gid={gid}"

    async def _fetch_raw_text(self, sheet_id: str, gid: str = "0") -> str:
        url = self._csv_url(sheet_id, gid)
        async with httpx.AsyncClient(follow_redirects=True, timeout=30.0) as client:
            response = await client.get(url)
            response.raise_for_status()
        return response.text

    def _parse_csv_text(self, text: str) -> List[Dict[str, Any]]:
        reader = csv.DictReader(io.StringIO(text))
        rows = []
        for row in reader:
            cleaned = {k.strip().strip('"'): v.strip().strip('"') for k, v in row.items() if k}
            rows.append(cleaned)
        return rows

    def _find_all_section_starts(self, lines: List[str]) -> List[int]:
        indices = []
        for i, line in enumerate(lines):
            if (line.startswith('"Name","Team"') or
                line.startswith('Name,Team') or
                line.startswith('"Name","Team","Nationality"')):
                indices.append(i)
        return indices

    def _parse_section_from_line(self, lines: List[str], start: int, end: Optional[int] = None) -> List[Dict]:
        section_lines = lines[start:end] if end else lines[start:]
        text = "\n".join(section_lines)
        rows = self._parse_csv_text(text)
        return [r for r in rows if r.get("Name", "").strip() not in ["", "Name"]]

    async def _fetch_sheets_section(self, year: int, section_index: int, core_cols: List[str]) -> List[Dict]:
        sheet_id = SHEET_URLS[year]
        text = await self._fetch_raw_text(sheet_id)
        lines = text.strip().split("\n")
        sections = self._find_all_section_starts(lines)

        if len(sections) <= section_index:
            return []

        start = sections[section_index]
        end = sections[section_index + 1] if len(sections) > section_index + 1 else None
        rows = self._parse_section_from_line(lines, start, end)
        rows = [r for r in rows if r.get("G", "0") not in ["", "0"]]

        result = []
        for row in rows:
            filtered = {k: v for k, v in row.items() if k in core_cols}
            result.append(filtered)
        return result

    # ─────────────────────────────────────────────
    # PUBLIC API
    # ─────────────────────────────────────────────

    async def fetch_hitting(self, year: int) -> List[Dict[str, Any]]:
        if year == 2025:
            return self._read_excel_sheet(2025, "Hitting Cumulative", HITTING_CORE_COLS)
        return await self._fetch_sheets_section(year, 0, HITTING_CORE_COLS)

    async def fetch_pitching(self, year: int) -> List[Dict[str, Any]]:
        if year == 2025:
            return self._read_excel_sheet(2025, "Pitching Cumulative", PITCHING_CORE_COLS)
        return await self._fetch_sheets_section(year, 1, PITCHING_CORE_COLS)

    async def fetch_fielding(self, year: int) -> List[Dict[str, Any]]:
        if year == 2025:
            return self._read_excel_sheet(2025, "Fielding Cumulative", FIELDING_CORE_COLS)
        return await self._fetch_sheets_section(year, 2, FIELDING_CORE_COLS)

    async def fetch_stats(self, year: int) -> List[Dict[str, Any]]:
        return await self.fetch_hitting(year)